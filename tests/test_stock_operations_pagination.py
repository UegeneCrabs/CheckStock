"""History pagination against an isolated DB; no app startup or external services."""

import html
import io
import re
import sqlite3
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook

with patch("dotenv.dotenv_values", return_value={}):
    from app.core.stores import STORES
    from app.repositories import operations as repo
    from app.web.routers import stock_operations as api


class OperationsPaginationTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:", check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.addCleanup(self.conn.close)
        self.conn.executescript(
            """
            CREATE TABLE stock_operations (
                id INTEGER PRIMARY KEY, store_slug TEXT, kind TEXT, source_type TEXT,
                source_name TEXT, sheet_url TEXT, from_fulfillment TEXT, from_marketplace TEXT,
                to_fulfillment TEXT, to_marketplace TEXT, note TEXT, user_name TEXT, created_at TEXT
            );
            CREATE TABLE stock_operation_items (
                id INTEGER PRIMARY KEY, operation_id INTEGER, article TEXT, barcode TEXT,
                name TEXT, quantity INTEGER
            );
            """
        )
        borrowed = SimpleNamespace(execute=self.conn.execute, close=lambda: None)
        self.allowed = tuple((slug, "WB") for slug in STORES)
        patches = (
            patch.object(repo, "get_connection", return_value=borrowed),
            patch.object(api, "scope_pairs", side_effect=lambda _user: self.allowed),
            patch.object(api, "has_action_permission", return_value=True),
            patch.object(
                api, "render_page", side_effect=lambda _title, _section, content, *_args, **_kw: content
            ),
            patch.object(api, "settings", SimpleNamespace(operation_history_limit=500)),
        )
        for item in patches:
            item.start()
            self.addCleanup(item.stop)
        app = FastAPI()

        @app.middleware("http")
        async def user_context(request: Request, call_next):
            request.state.user = object()
            return await call_next(request)

        app.include_router(api.router)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def add_operation(self, slug="tris", kind="delivery", source=None, target="WB", name="Сотрудник"):
        row = self.conn.execute(
            """INSERT INTO stock_operations
               (store_slug, kind, source_type, from_marketplace, to_marketplace, user_name, created_at)
               VALUES (?, ?, 'manual', ?, ?, ?, '2026-08-07T07:20:00+00:00')""",
            (slug, kind, source, target, name),
        )
        operation_id = row.lastrowid
        self.conn.execute(
            "INSERT INTO stock_operation_items (operation_id, article, name, quantity) VALUES (?, ?, ?, ?)",
            (operation_id, str(operation_id), "Товар", -2 if kind == "shipment" else 2),
        )
        return operation_id

    @staticmethod
    def displayed_ids(response):
        return [int(value) for value in re.findall(r"/admin/operations/(\d+)/xlsx", response.text)]

    def test_all_stores_can_reach_history_beyond_500(self):
        for slug in STORES:
            with self.subTest(store=slug):
                ids = [self.add_operation(slug) for _ in range(632)]
                first = self.client.get(f"/stock/{slug}/operations")
                second = self.client.get(f"/stock/{slug}/operations?page=2")
                self.assertEqual(first.status_code, 200)
                self.assertEqual(second.status_code, 200)
                self.assertEqual(self.displayed_ids(first), list(reversed(ids))[:500])
                self.assertEqual(self.displayed_ids(second), list(reversed(ids))[500:])
                self.assertIn("Записи 501–632 из 632", html.unescape(re.sub(r"<[^>]+>", "", second.text)))
                self.assertIn("Страница 2 из 2", second.text)
                self.assertIn("<span>Операций</span><strong>632</strong>", second.text)
                self.assertIn(f"/stock/{slug}/operations?kind=&amp;page=2", first.text)

    def test_scope_applies_before_pagination_and_to_totals(self):
        ids = [self.add_operation() for _ in range(502)]
        ids.append(self.add_operation(kind="transfer", source="OZON", target="WB"))
        ids.append(self.add_operation(target=None))
        for _ in range(510):
            self.add_operation(target="OZON", name="Скрытый сотрудник")
            self.add_operation("rimili")
        self.allowed = (("tris", "WB"),)
        response = self.client.get("/stock/tris/operations?page=2")
        self.assertEqual(self.displayed_ids(response), list(reversed(ids))[500:])
        self.assertIn("Записи 501–504 из 504", html.unescape(re.sub(r"<[^>]+>", "", response.text)))
        stats = repo.get_store_operation_stats("tris", marketplaces=("WB",))
        self.assertEqual(stats["summary"], {"total": 504, "positions": 504, "units": 1008, "employees": 1})
        self.assertEqual(sum(stats["counts"].values()), 504)
        self.allowed = ()
        empty = self.client.get("/stock/tris/operations")
        self.assertEqual(self.displayed_ids(empty), [])
        self.assertIn("Пока нет операций", empty.text)

    def test_kind_filter_keeps_all_transfer_stages_and_navigation(self):
        ids = [self.add_operation(kind=api.TRANSFER_KINDS[n % 5]) for n in range(503)]
        for _ in range(3):
            self.add_operation(kind="shipment")
        response = self.client.get("/stock/tris/operations?kind=transfer&page=2")
        self.assertEqual(self.displayed_ids(response), list(reversed(ids))[500:])
        self.assertIn("Записи 501–503 из 503", html.unescape(re.sub(r"<[^>]+>", "", response.text)))
        self.assertIn('href="/stock/tris/operations?kind=transfer&amp;page=1"', response.text)
        self.assertIn('Перемещения<span class="ops-filter-count">503</span>', response.text)
        self.assertIn('Все<span class="ops-filter-count">506</span>', response.text)
        self.assertIn("<span>Операций</span><strong>503</strong>", response.text)
        stats = repo.get_store_operation_stats("tris", ("shipment",), marketplaces=("WB",))
        self.assertEqual(stats["summary"]["units"], 6)

    def test_empty_out_of_range_and_invalid_pages(self):
        empty = self.client.get("/stock/tris/operations?page=100")
        self.assertIn("Движений пока не было", empty.text)
        self.assertIn("Страница 1 из 1", empty.text)
        operation_id = self.add_operation()
        response = self.client.get("/stock/tris/operations?page=999999999999999999999999999")
        self.assertEqual(self.displayed_ids(response), [operation_id])
        self.assertIn("Страница 1 из 1", response.text)
        for page in ("0", "-1", "invalid"):
            self.assertEqual(self.client.get(f"/stock/tris/operations?page={page}").status_code, 422)
        self.assertEqual(self.client.get("/stock/missing/operations").status_code, 404)

    def test_excel_contains_all_filtered_accessible_operations_and_items(self):
        ids = [self.add_operation(kind="transfer_dispatch") for _ in range(632)]
        self.add_operation(kind="shipment")
        self.add_operation(kind="transfer_dispatch", target="OZON")
        self.add_operation("rimili", kind="transfer_dispatch")
        response = self.client.get("/stock/tris/operations/xlsx?kind=transfer")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["Операции"].max_row, 633)
        self.assertEqual(workbook["Позиции"].max_row, 633)
        article_ids = [int(row[5]) for row in workbook["Позиции"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(article_ids, list(reversed(ids)))
        self.allowed = ()
        self.assertEqual(self.client.get("/stock/tris/operations/xlsx").status_code, 403)


if __name__ == "__main__":
    unittest.main()
