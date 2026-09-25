"""F05 regression checks on synthetic inputs and a temporary database; no secrets/network."""

import asyncio
import importlib
import io
import json
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import Request
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql, sqlite
from sqlalchemy.schema import CreateTable

WB = {
    "retail_price": 1000, "customer_price": 900, "purchase_price": 200, "fulfillment_cost": 30,
    "subject_commission_percent": 10, "wb_extra_tariff_percent": 1, "acquiring_percent": 2,
    "delivery_wb_rub": 50, "return_cost_rub": 40, "buyout_percent": 80,
    "volume_l": 2, "acceptance_coefficient": 1, "storage_wb_rub": 2, "turnover_days": 10,
    "team_commission_percent": 3, "tax_system": "usn", "vat_percent": 20, "usn_percent": 6,
    "orders_count": 10, "advertising_spend": 100,
}
YM = {
    "seller_price": 1000, "buyer_price": 900, "purchase_price": 200, "fulfillment_cost": 30,
    "commission_percent": 10, "payment_acceptance": 2, "acquiring_percent": 1.6,
    "delivery_cost": 50, "middle_mile": 20, "transit_cost": 5, "buyout_percent": 80,
    "company_commission_percent": 3, "vat_percent": 20, "usn_percent": 6,
    "loss_percent": 1, "disposal_cost": 10, "orders_count": 10, "advertising_spend": 100,
}
DAY = "2026-09-23"


class CompletenessTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with patch("dotenv.dotenv_values", return_value={}), patch.dict(os.environ, {}, clear=True):
            cls.daily = importlib.import_module("app.economics.daily_calculation")
            cls.repo = importlib.import_module("app.repositories.daily_economics")
            cls.core = importlib.import_module("app.repositories.core")
            cls.database_module = importlib.import_module("app.infrastructure.database")
            cls.orm = importlib.import_module("app.infrastructure.orm")
            importlib.import_module("app.repositories.schema")
            cls.routes = importlib.import_module("app.web.routers.unit_economics")
            cls.wb_history = importlib.import_module("app.economics.wb.history")
            cls.ym_calc = importlib.import_module("app.yandex.economics_calculation")
            cls.ym_days = importlib.import_module("app.yandex.economics_days")
            cls.ym_reports = importlib.import_module("app.economics.yandex.reports")
            cls.reporting = importlib.import_module("app.economics.reporting")
            cls.wb_export = importlib.import_module("app.economics.wb.report_export")
            cls.ym_export = importlib.import_module("app.economics.yandex.report_export")
            cls.coverage = importlib.import_module("app.repositories.economics_coverage")
            cls.wb_repo = importlib.import_module("app.repositories.unit_economics_1c")
            cls.wb_funnel = importlib.import_module("app.wb.funnel_orders")

    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix="checkstock-f05-")
        self.addCleanup(temp.cleanup)
        self.database = self.database_module.Database(Path(temp.name) / "f05.sqlite3")
        self.addCleanup(self.database.dispose)
        self.orm.OrmBase.metadata.create_all(self.database.engine)
        redirect = patch.object(self.core, "database_for_path", return_value=self.database)
        redirect.start()
        self.addCleanup(redirect.stop)

    def calculate(self, market="WB", values=None):
        return self.daily.calculate(market, values if values is not None else WB)[1]

    def wb_snapshot(self, values):
        return {"inputs_json": json.dumps(values), "calculation_version": 3,
                "unit_margin": self.calculate(values=values)["margin"], "purchase_price": values.get("purchase_price")}

    def period(self, snapshots, orders, ads, orders_days, ad_days, end=DAY):
        return self.routes._report_historical_economics(
            date_from=date.fromisoformat(DAY), date_to=date.fromisoformat(end), daily_orders=orders,
            margin_snapshots=snapshots, daily_advertising=ads, orders_days=orders_days, advertising_days=ad_days,
            live_day=date(2026, 9, 25), live_unit_margin=999999, live_purchase_price=999,
        )

    def test_complete_data_preserve_formulas(self):
        wb = self.calculate()
        ym = self.calculate("YANDEX MARKET", YM)
        self.assertEqual((wb["margin"], wb["day_profit"], wb["roi"]), (323.6, 2488.8, 161.8))
        self.assertEqual((ym["margin"], ym["day_profit"], ym["roi"]), (352.9, 2723.2, 176.45))
        self.assertTrue(wb["daily_complete"])
        self.assertTrue(ym["daily_complete"])

    def test_unknown_advertising_is_partial_real_zero_is_complete(self):
        for market, fixture in (("WB", WB), ("YANDEX MARKET", YM)):
            with self.subTest(market=market):
                unknown = self.calculate(market, {**fixture, "advertising_spend": None})
                zero = self.calculate(market, {**fixture, "advertising_spend": 0})
                self.assertEqual(unknown["day_profit"], zero["day_profit"])
                self.assertFalse(unknown["daily_complete"])
                self.assertTrue(zero["daily_complete"])
                self.assertIn("advertising_spend", unknown["daily_missing"])
                self.assertEqual(unknown["purchase_value"], 1600)

    def test_missing_purchase_and_multiple_costs_keep_partial_margin(self):
        for market, fixture in (("WB", WB), ("YANDEX MARKET", YM)):
            with self.subTest(market=market):
                values, result = self.daily.calculate(market, {**fixture, "purchase_price": None, "fulfillment_cost": None})
                self.assertGreater(result["margin"], self.calculate(market, fixture)["margin"])
                self.assertIsNone(result["roi"])
                self.assertIsNone(result["purchase_value"])
                self.assertIsNone(values["purchase_price"])
                self.assertIn("purchase_price", result["missing"])
                self.assertIn("fulfillment_cost", result["missing"])
                self.assertEqual(result["status"], "Неполный расчёт")

    def test_unknown_vat_does_not_invent_usn_base(self):
        wb = self.calculate(values={**WB, "vat_percent": None})
        ym = self.calculate("YANDEX MARKET", {**YM, "vat_percent": None})
        self.assertEqual(wb["margin"], 518.6)
        self.assertIsNone(wb["usn"])
        self.assertEqual(ym["margin"], 547.9)
        self.assertIsNone(ym["costs"]["usn"])

    def test_missing_price_prevents_unit_and_positive_order_day(self):
        for market, fixture, price in (("WB", WB, "retail_price"), ("YANDEX MARKET", YM, "seller_price")):
            for value in (None, 0):
                with self.subTest(market=market, price=value):
                    result = self.calculate(market, {**fixture, price: value})
                    self.assertIsNone(result["margin"])
                    self.assertIsNone(result["day_profit"])
                    self.assertEqual(result["status"], "Недостаточно данных")
                    self.assertEqual(result["purchase_value"], 1600)

    def test_zero_purchase_does_not_produce_zero_roi(self):
        for market, fixture in (("WB", WB), ("YANDEX MARKET", YM)):
            result = self.calculate(market, {**fixture, "purchase_price": 0})
            self.assertIsNone(result["roi"])
            self.assertNotIn("purchase_price", result["missing"])

    def test_wb_snapshot_does_not_treat_dto_defaults_or_stale_prices_as_observations(self):
        from app.dto.unit_economics_1c import UnitEconomics1CCabinetSettings, UnitEconomics1CProductSettings

        args = {"snapshot_day": date.fromisoformat(DAY), "store_slug": "rimili", "article": "123",
                "price_snapshot": {"day": DAY, "retail_price": 1000, "customer_price_with_spp": 900},
                "product_metrics": self.wb_history.unit_economics_1c.empty_product_metrics(today=date.fromisoformat(DAY)),
                "product_settings": UnitEconomics1CProductSettings(store_slug="rimili", article="123"),
                "product_reference": {}, "cabinet": UnitEconomics1CCabinetSettings(store_slug="rimili"),
                "captured_at": DAY + "T23:00:00Z"}
        values = json.loads(self.wb_history.calculate_snapshot_row(**args)["inputs_json"])
        for field in ("purchase_price", "fulfillment_cost", "buyout_percent", "orders_count", "advertising_spend", "acquiring_percent", "volume_l"):
            self.assertIsNone(values[field], field)
        self.assertEqual(values["retail_price"], 1000)
        args["cabinet"] = UnitEconomics1CCabinetSettings(store_slug="rimili", default_buyout_percent=90, updated_at=DAY)
        values = json.loads(self.wb_history.calculate_snapshot_row(**args)["inputs_json"])
        self.assertEqual(values["buyout_percent"], 90)
        args["price_snapshot"] = {**args["price_snapshot"], "day": "2026-09-22"}
        values = json.loads(self.wb_history.calculate_snapshot_row(**args)["inputs_json"])
        self.assertIsNone(values["retail_price"])
        self.assertIsNone(values["customer_price"])
        self.assertIn("2026-09-22", values["price_missing_reason"])

    def test_zero_orders_retain_actual_advertising_even_without_price(self):
        for market, fixture, price in (("WB", WB, "retail_price"), ("YANDEX MARKET", YM, "seller_price")):
            result = self.calculate(market, {**fixture, price: None, "orders_count": 0, "advertising_spend": 123})
            self.assertEqual(result["day_profit"], -123)
            self.assertEqual(result["purchase_value"], 0)
            self.assertTrue(result["daily_complete"])
            unknown = self.calculate(market, {**fixture, "orders_count": None, "advertising_spend": 123})
            self.assertIsNone(unknown["day_profit"])
            self.assertIn("orders_count", unknown["daily_missing"])

    def test_wb_period_one_partial_day_and_separate_purchase(self):
        next_day = "2026-09-24"
        result = self.period({d: self.wb_snapshot(WB) for d in (DAY, next_day)},
                             {d: {"orders_count": 10} for d in (DAY, next_day)},
                             {DAY: 100, next_day: 999999}, {DAY, next_day}, {DAY}, end=next_day)
        self.assertEqual(result["margin"], 5077.6)
        self.assertEqual(result["purchase_value"], 3200)
        self.assertEqual(result["missing_parameters"], {next_day: ["advertising_spend"]})
        self.assertFalse(result["complete"])
        self.assertIsNotNone(result["roi"])
        self.assertIn(next_day, " ".join(result["messages"]))

    def test_wb_missing_source_not_inferred_from_numeric_dictionary(self):
        result = self.period({DAY: self.wb_snapshot(WB)}, {DAY: {"orders_count": 10}}, {DAY: 0}, set(), set())
        self.assertIsNone(result["margin"])
        self.assertIsNone(result["advertising_spend"])
        self.assertIn("orders_count", result["missing_parameters"][DAY])

    def test_wb_no_current_values_in_missing_past_day(self):
        result = self.period({}, {DAY: {"orders_count": 10}}, {DAY: 100}, {DAY}, {DAY})
        self.assertIsNone(result["margin"])
        self.assertEqual(result["unavailable_days"], [DAY])

    def test_unknown_advertising_does_not_hide_day_purchase_or_price_in_details(self):
        item = self.routes._report_daily_calculations(
            date_from=date.fromisoformat(DAY), date_to=date.fromisoformat(DAY),
            daily_orders={DAY: {"orders_count": 10}}, margin_snapshots={DAY: self.wb_snapshot(WB)},
            live_day=date(2026, 9, 25), live_snapshot=None, daily_advertising={}, orders_days={DAY}, advertising_days=set(),
        )[0]
        self.assertEqual(item["retail_price"], 1000)
        self.assertEqual(item["purchase_price"], 200)
        self.assertEqual(item["net_profit"], 323.6)
        self.assertEqual(item["status"], "Неполный расчёт")

    def test_yandex_report_day_and_period_share_partial_result(self):
        saved = {"day": DAY, "inputs": YM, "calculation_version": 15}
        resolved = self.ym_days.resolve_day(DAY, saved, {"orders_count": 10}, {"spend": 7777}, True, False)
        detail = self.ym_reports.daily_row(DAY, resolved, {"orders_count": 10}, {}, True, False, 5)
        period = self.ym_calc.aggregate([resolved], [DAY])
        self.assertEqual(period["margin"], 2823.2)
        self.assertEqual(detail["day_profit"], period["margin"])
        self.assertEqual(detail["missing"], ["advertising_spend"])
        self.assertEqual(detail["retail_price"], 1000)
        self.assertFalse(period["complete"])

    def test_aggregate_explains_uncalculable_product_and_matches_roi_scope(self):
        rows = [{"article": "A", "margin": 100, "purchase_value": 200, "roi_purchase_value": 200,
                 "margin_complete": True},
                {"article": "B", "margin": None, "purchase_value": 700, "margin_complete": False,
                 "messages": ["Нет цены за " + DAY]}]
        totals = self.reporting._unit_profit_report_totals(rows)
        self.assertEqual(totals["margin"], 100)
        self.assertEqual(totals["purchase_value"], 900)
        self.assertEqual(totals["roi"], 50)
        self.assertFalse(totals["margin_complete"])
        self.assertEqual(totals["covered_products"], 1)
        self.assertIn("1 из 2", " ".join(totals["messages"]))
        rows[0]["roi_purchase_value"] = None
        self.assertIsNone(self.reporting._unit_profit_report_totals(rows)["roi"])

    def test_successful_empty_wb_sources_mark_zero_and_refresh_only_dated_metrics(self):
        key = ("WB", "rimili", "123", DAY)
        self.repo.capture(key, self.repo.observation({**WB, "orders_count": None, "advertising_spend": None}))
        self.wb_funnel._replace_day("rimili", date.fromisoformat(DAY), [])
        self.wb_repo.replace_daily_advertising("rimili", DAY, DAY, [])
        data = self.repo.get(key)
        self.assertEqual(data["values"]["orders_count"], 0)
        self.assertEqual(data["values"]["advertising_spend"], 0)
        self.assertEqual(data["values"]["retail_price"], 1000)
        self.assertEqual(data["result"]["day_profit"], 0)
        self.assertTrue(data["result"]["daily_complete"])
        self.assertEqual(self.coverage.wb_days(("rimili",), DAY, DAY)["rimili"], {"orders": {DAY}, "advertising": {DAY}})

    def test_marker_rolls_back_with_source_and_is_store_scoped(self):
        with self.database.connect() as conn:
            self.coverage.mark_loaded(conn, "rimili", "advertising", DAY, DAY)
            conn.rollback()
        self.assertFalse(self.coverage.wb_days(("rimili",), DAY, DAY)["rimili"]["advertising"])
        self.wb_repo.replace_daily_advertising("rimili", DAY, DAY, [])
        self.assertFalse(self.coverage.wb_days(("gogol",), DAY, DAY)["gogol"]["advertising"])

    def test_manual_fill_recalculates_and_source_refresh_preserves_override_and_other_day(self):
        for market, fixture in (("WB", WB), ("YANDEX MARKET", YM)):
            key = (market, "rimili", "123", DAY)
            other = (*key[:3], "2026-09-22")
            source = self.repo.observation({**fixture, "fulfillment_cost": None, "advertising_spend": None})
            self.repo.capture(key, source)
            self.repo.capture(other, source)
            before_other = self.repo.get(other)
            data = self.repo.get(key)
            preview, _ = self.repo.proposal(key, data, "fulfillment_cost", 30)
            self.repo.correct(key, token=data["token"], preview_token=preview["preview_token"],
                              field="fulfillment_cost", value=30, reason="Synthetic correction", actor="Test")
            self.repo.capture(key, self.repo.observation({**fixture, "fulfillment_cost": 99}))
            after = self.repo.get(key)
            self.assertEqual(after["values"]["fulfillment_cost"], 30)
            self.assertTrue(after["result"]["daily_complete"])
            self.assertEqual(after["result"]["day_profit"], self.calculate(market, fixture)["day_profit"])
            self.assertEqual(self.repo.get(other), before_other)
            self.assertIsNone(after["original"]["values"]["fulfillment_cost"])

    def test_manual_orders_and_ads_take_precedence_in_both_reports(self):
        overrides = {"orders_count": 0, "advertising_spend": 73}
        wb = self.wb_history.report_day(DAY, {**self.wb_snapshot(WB), "overrides": overrides}, {}, None)
        ym = self.ym_days.resolve_day(DAY, {"inputs": YM, "overrides": overrides}, {}, {}, False, False)
        self.assertEqual(wb["profit"], -73)
        self.assertEqual(ym["profit"], -73)
        self.assertTrue(wb["complete"])
        self.assertTrue(ym["complete"])

    def test_sqlite_and_postgres_schema_compile(self):
        for dialect in (sqlite.dialect(), postgresql.dialect()):
            for table in ("economics_source_days", "economics_daily", "economics_daily_events"):
                sql = str(CreateTable(self.orm.OrmBase.metadata.tables[table]).compile(dialect=dialect))
                self.assertIn("PRIMARY KEY", sql)

    def test_wb_actual_report_builder_and_export_with_missing_ads(self):
        from app.dto.identity import User

        with self.database.connect() as conn:
            conn.execute("INSERT INTO stock_items (store_slug,marketplace,article,barcode,name) VALUES ('rimili','WB','123','111','F05 synthetic')")
            conn.commit()
        self.wb_funnel._replace_day("rimili", date.fromisoformat(DAY), [("123", "sku", "F05 synthetic", 10, 9000, 0, 0, 8, 7200, 80)])
        self.repo.capture(("WB", "rimili", "123", DAY), self.repo.observation({**WB, "advertising_spend": None}, version=3))
        user = User(id=1, login="f05", full_name="F05", role="superadmin", created_at="2026-09-01T00:00:00Z")
        request = Request({"type": "http", "headers": [], "query_string": f"date_from={DAY}&date_to={DAY}&daily_details=1".encode(), "state": {"user": user}})
        with patch.object(self.routes, "accessible_stores", return_value=("rimili",)), patch.object(self.routes.reports_cache, "get", side_effect=lambda _key, loader: loader()):
            screen = asyncio.run(self.routes._unit_economics_1c_unit_profit_report_data(request))
            export = asyncio.run(self.routes._unit_economics_1c_unit_profit_report_data(request, for_export=True))
        self.assertTrue(screen["ok"])
        self.assertEqual(len(screen["rows"]), 1)
        row = screen["rows"][0]
        self.assertEqual(row["margin"], 2588.8)
        self.assertFalse(row["margin_complete"])
        self.assertEqual(row["missing_parameters"], {DAY: ["advertising_spend"]})
        self.assertIsNone(row["advertising_spend"])
        self.assertEqual(row["daily_calculations"][0]["day_profit"], row["margin"])
        self.assertEqual(export["rows"][0]["margin"], row["margin"])
        self.assertEqual(export["totals"]["messages"], screen["totals"]["messages"])
        content, _ = self.wb_export.build_xlsx(export)
        self.assertTrue(content.startswith(b"PK"))

        # Read the real main-table endpoint against the same saved day. A page
        # read may recalculate locally but must not capture/backfill observations.
        request = Request({"type": "http", "headers": [], "query_string": f"data=1&store=rimili&article=123&date_from={DAY}&date_to={DAY}".encode(), "state": {"user": user}})
        with patch.object(self.routes, "accessible_stores", return_value=("rimili",)), patch.object(self.repo, "capture", side_effect=AssertionError("Read attempted a write")), patch.object(
            self.routes.stock_sheet_inbound,
            "load",
            side_effect=lambda _store, _marketplace, catalog: SimpleNamespace(
                catalog=catalog, quantities={}, confirmed_quantities={}, available=True
            ),
        ):
            payload = json.loads(asyncio.run(self.routes.sales_unit_economics_1c(request)).body)
        self.assertEqual(payload["product"]["economics_7d"]["margin"], row["margin"])
        self.assertEqual(payload["product"]["economics_7d"]["roi"], row["roi"])
        self.assertEqual(payload["product"]["economics_7d"]["messages"], row["messages"])
        compact = self.routes._unit_economics_1c_product_summary(payload["product"])
        self.assertFalse(compact["economics_7d"]["complete"])
        self.assertIn("daily_complete", compact["current_economics"])

    def test_yandex_actual_report_matches_history_and_preserves_past_buyout(self):
        market = "YANDEX MARKET"
        next_day = "2026-09-24"
        for day in (DAY, next_day):
            self.repo.capture((market, "rimili", "123", day), self.repo.observation(YM, version=15))
        orders = [{"article": "123", "day": day, "orders_count": 10, "orders_amount": 9000} for day in (DAY, next_day)]
        self.ym_reports.metrics.save_daily("rimili", "orders", orders, DAY, next_day, DAY + "T23:00:00Z")
        self.ym_reports.metrics.save_daily("rimili", "advertising", [{"article": "123", "day": DAY, "spend": 100}], DAY, DAY, DAY + "T23:00:00Z")
        product = {"store_slug": "rimili", "article": "123", "stock": {"total": 1, "fbs": 1, "fbo": 0, "fulfillment": 0, "days": None, "average_daily_orders": None}, "advertising": {}}
        with patch.object(self.ym_reports, "catalog", return_value=[{"store_slug": "rimili", "article": "123", "name": "Synthetic"}]), patch.object(self.ym_reports, "load_products", return_value=[product]), patch.object(self.ym_reports.economics, "context", return_value={}), patch.object(self.ym_reports.economics, "effective", return_value={"values": {**YM, "buyout_percent": 10}}), patch.object(self.repo, "capture", side_effect=AssertionError("Read attempted a write")):
            row = self.ym_reports.load_rows(("rimili",), None, date.fromisoformat(DAY), date.fromisoformat(next_day), today=date(2026, 9, 25), include_details=True)[0]
        self.assertEqual(row["margin"], 5546.4)
        self.assertEqual(row["purchase_value"], 3200)
        self.assertEqual(row["buyout_percent"], 80)
        self.assertEqual(row["expected_buyout_amount"], 14400)
        self.assertEqual(row["missing_parameters"], {next_day: ["advertising_spend"]})
        self.assertEqual(sum(d["day_profit"] for d in row["daily_calculations"]), row["margin"])
        economics = self.ym_reports.economics
        history = economics.resolved_history("rimili", "123", economics.repository.history("rimili", DAY, next_day), [DAY, next_day])
        self.assertEqual(economics.aggregate(history, [DAY, next_day])["margin"], row["margin"])

    def test_xlsx_unknown_buyouts_stay_empty_and_roi_basis_is_numeric(self):
        self.assertIsNone(self.wb_export._cell_value({"orders_count": None, "buyout_percent": 80}, "expected_buyouts", "number"))
        self.assertEqual(self.wb_export._cell_value({"orders_count": 0, "buyout_percent": None}, "expected_buyouts", "number"), 0)
        self.assertIsNone(self.wb_export._cell_value({"expected_buyouts": None}, "expected_buyouts", "number"))
        self.assertEqual(self.wb_export._cell_value({"expected_buyouts": 8}, "expected_buyouts", "number"), 8)

    def test_full_buyout_does_not_require_return_tariff(self):
        result = self.ym_calc.calculate({**YM, "buyout_percent": 100, "return_cost": None}, without_advertising=True)
        self.assertNotIn("return_cost", result["missing"])
        self.assertEqual(result["logistics"]["returns"], 0)
        self.assertIsNotNone(result["logistics"]["total"])

    def test_group_totals_preserve_unknown_sources_and_explain_omissions(self):
        partial = {"date": DAY, "orders_count": 10, "expected_buyouts": 8, "buyout_percent": 80,
                   "net_profit": 100, "day_profit": 800, "available": True, "complete": False,
                   "missing": ["advertising_spend"]}
        unavailable = {"date": DAY, "orders_count": None, "expected_buyouts": None,
                       "net_profit": 99999, "day_profit": None, "available": False, "complete": False}
        rows = [{"daily_calculations": [partial]}, {"daily_calculations": [unavailable]}]
        total = self.reporting._aggregate_report_daily_calculations(rows)[0]
        self.assertEqual(total["net_profit"], 100)
        self.assertEqual(total["day_profit"], 800)
        self.assertEqual(total["buyout_percent"], 80)
        self.assertIsNone(total["advertising_spend"])
        self.assertIsNone(total["advertising_per_unit"])
        self.assertIn("1 из 2", " ".join(total["messages"]))
        summary = self.reporting._unit_profit_report_totals([
            {"margin": 800, "margin_complete": False, "margin_missing_days": [DAY], "unavailable_days": []},
        ])
        self.assertEqual(summary["unavailable_days"], [])
        self.assertIsNone(summary["drr"])
        self.assertIsNone(summary["cpc"])

    def test_xlsx_numeric_partial_values_have_status_and_reasons(self):
        row = {"name": "Synthetic", "article": "123", "margin": 10000, "purchase_value": 500,
               "roi_purchase_value": 500,
               "roi": 2000, "margin_complete": False, "status": "Неполный расчёт",
               "messages": ["Не учтены: расходы на рекламу за " + DAY], "margin_missing_days": [DAY],
               "daily_calculations": [{"date": DAY, "net_profit": 100, "day_profit": 10000,
                    "status": "Неполный расчёт", "messages": ["Расходы на рекламу за " + DAY]}]}
        report = {"rows": [row], "totals": row, "period_from": DAY, "period_to": DAY}
        for exporter, header_row, data_row in ((self.wb_export, 5, 6), (self.ym_export, 2, 4)):
            content, _ = exporter.build_xlsx(report)
            book = load_workbook(io.BytesIO(content))
            sheet = book.active
            columns = {cell.value: cell.column for cell in sheet[header_row]}
            self.assertEqual(sheet.cell(data_row, columns["Маржа периода, ₽"]).value, 10000)
            self.assertEqual(sheet.cell(data_row, columns["Маржа периода, ₽"]).data_type, "n")
            self.assertEqual(sheet.cell(data_row, columns["Закупка в базе ROI, ₽"]).value, 500)
            self.assertEqual(sheet.cell(data_row, columns["Полнота расчёта"]).value, "Неполный расчёт")
            self.assertTrue(any(DAY in str(c.value) and "реклам" in str(c.value) for c in sheet[data_row]))
            book.close()


if __name__ == "__main__":
    unittest.main()
