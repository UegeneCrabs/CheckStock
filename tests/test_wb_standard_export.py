import unittest
from datetime import date
from io import BytesIO
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook

from app.dto.wb_standard_costs import Reports, StorageCalculation
from app.web.routers.wb_standard_costs import ExportRequest, export_rows


class ExportTests(unittest.TestCase):
    def payload(self, articles):
        return ExportRequest(
            rows=[
                {
                    "store_slug": "tris",
                    "article": str(article),
                    "scheme": "FBW",
                    "period_from": "2026-10-12",
                    "period_to": "2026-10-18",
                }
                for article in articles
            ]
        )

    def reports(self):
        return Reports.model_validate(
            {
                "stores": [
                    {
                        "slug": "tris",
                        "name": "TRIS",
                        "report": {
                            "date_from": "2026-10-12",
                            "date_to": "2026-10-18",
                            "rows": [
                                {
                                    "article": str(i),
                                    "scheme": "FBW",
                                    "name": "=1+1",
                                    "profit": {"margin": -12.5, "roi": -2, "missing": ["Закупка"]},
                                    "current_profit": {"margin": None},
                                    "delta": {"margin": 3},
                                }
                                for i in range(60)
                            ],
                        },
                    }
                ]
            }
        )

    def test_all_selected_rows_order_numbers_missing_values_and_safe_text(self):
        with patch("app.web.routers.wb_standard_costs.load_reports", return_value=self.reports()):
            content = export_rows(self.payload(range(59, -1, -1)), {"tris"})
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet.max_row, 61)
        self.assertEqual(sheet["B2"].value, "59")
        self.assertEqual(sheet["B61"].value, "0")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertEqual(sheet["C2"].value, "=1+1")
        self.assertEqual(sheet["K2"].value, -12.5)
        self.assertEqual(sheet["K2"].number_format, "#,##0")
        self.assertEqual(sheet["K2"].font.color.rgb, "00C62828")
        self.assertIsNone(sheet["M2"].value)
        self.assertEqual(sheet["O2"].font.color.rgb, "00188038")
        self.assertEqual(sheet.freeze_panes, "C2")

    def test_forbidden_store_is_not_loaded(self):
        with patch("app.web.routers.wb_standard_costs.load_reports") as loader:
            with self.assertRaises(HTTPException) as error:
                export_rows(self.payload([1]), {"rimili"})
        self.assertEqual(error.exception.status_code, 403)
        loader.assert_not_called()

    def test_stale_selection_fails_instead_of_partial_export(self):
        with patch("app.web.routers.wb_standard_costs.load_reports", return_value=self.reports()):
            with self.assertRaises(HTTPException) as error:
                export_rows(self.payload([99]), {"tris"})
        self.assertEqual(error.exception.status_code, 409)

    def test_changed_period_is_rejected_even_if_article_still_exists(self):
        payload = self.payload([1])
        payload = ExportRequest(rows=[payload.rows[0].model_copy(update={"period_to": date(2026, 10, 19)})])
        with patch("app.web.routers.wb_standard_costs.load_reports", return_value=self.reports()):
            with self.assertRaises(HTTPException) as error:
                export_rows(payload, {"tris"})
        self.assertEqual(error.exception.status_code, 409)

    def test_assumed_storage_zero_preserves_numeric_value_and_warning(self):
        reports = self.reports()
        row = reports["stores"][0]["report"]["rows"][0]
        row = row.model_copy(update={"profit": row.profit.model_copy(update={"storage_total": 0})})
        row = row.model_copy(update={"storage_calculation": StorageCalculation(warning="Нет начислений API")})
        report = reports.stores[0].report
        report = report.model_copy(update={"rows": (row,) + report.rows[1:]})
        reports = reports.model_copy(
            update={"stores": (reports.stores[0].model_copy(update={"report": report}),)}
        )
        with patch("app.web.routers.wb_standard_costs.load_reports", return_value=reports):
            content = export_rows(self.payload([0]), {"tris"})
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet.cell(2, 27).value, 0)
        self.assertIn("†", sheet.cell(2, 27).number_format)
        self.assertEqual(sheet.cell(2, 27).comment.text, "Нет начислений API")
