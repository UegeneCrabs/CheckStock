from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.config import settings
from app.core.constants import OZON_REPORTS, OzonReportGroup
from app.dto.ozon_reports import OzonExportColumn, OzonProfitReport, OzonReportColumns, OzonTargetReport


def cell(value: object) -> str | int | float | bool | None:
    if isinstance(value, str):
        return "'" + value if value.startswith(("=", "+", "-", "@")) else value
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value)
    return str(value)


def build_workbook(report: OzonProfitReport | OzonTargetReport) -> bytes:
    columns = OzonReportColumns.model_validate_json(
        (settings.base_dir / OZON_REPORTS.columns_path).read_text(encoding="utf-8")
    )
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(title: str, fields: tuple[OzonExportColumn, ...], rows: list[dict[str, object]]) -> None:
        page = workbook.create_sheet(title)
        page.append([field.label for field in fields])
        for field in page[1]:
            field.font = Font(bold=True)
        for row in rows:
            page.append([cell(row.get(field.key)) for field in fields])
        page.freeze_panes = "D2"
        page.auto_filter.ref = page.dimensions
        for index, field in enumerate(fields, 1):
            page.column_dimensions[get_column_letter(index)].width = max(len(field.label), len(field.key))

    if isinstance(report, OzonTargetReport):
        sheet(columns.target_title, columns.target, [row.model_dump(mode="json") for row in report.rows])
    else:
        rows = report.category_rows if report.group_by == OzonReportGroup.SUBJECT else report.rows
        sheet(
            columns.profit_title,
            columns.profit,
            [
                {"name": columns.totals_label, **report.totals.model_dump(mode="json")},
                *[row.model_dump(mode="json") for row in rows],
            ],
        )
        identities = tuple(
            field for field in columns.profit if field.key in ("name", "article", "store_name")
        )
        fields = (
            OzonExportColumn(key="date", label=columns.daily_title, format="text"),
            *identities,
            *columns.daily,
            *columns.daily_diagnostics,
        )
        sheet(
            columns.daily_title,
            fields,
            [
                {
                    "name": row.name,
                    "article": row.article,
                    "store_name": row.store_name,
                    **day.model_dump(mode="json"),
                }
                for row in report.rows
                for day in row.daily_calculations
            ],
        )
    data = BytesIO()
    workbook.save(data)
    return data.getvalue()
