"""Excel presentation of the standard conditions report."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.dto.wb_standard_costs import ExportItem

COSTS = (
    ("retail", "Цена"),
    ("purchase", "Закупка"),
    ("fulfillment", "ФФ"),
    ("advertising", "Реклама"),
    ("acquiring", "Эквайринг"),
    ("team", "Комиссия компании"),
    ("vat", "НДС"),
    ("tax", "УСН / ОСНО"),
    ("commission", "Комиссия WB"),
    ("logistics", "Логистика"),
    ("storage_total", "Хранение"),
    ("total_cost", "Все расходы"),
)


def build_xlsx(rows: list[ExportItem]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Сравнение ИУ"
    headers = [
        "Магазин",
        "Артикул",
        "Название",
        "Схема",
        "Предмет",
        "Менеджер",
        "Период с",
        "Период по",
        "ТО, ₽",
        "ТО: полный период",
        "Без ИУ: чистая прибыль, ₽",
        "Без ИУ: ROI, %",
        "С ИУ: чистая прибыль, ₽",
        "С ИУ: ROI, %",
        "Разница прибыли, ₽",
        "Разница ROI, п.п.",
    ]
    headers += [f"{label}: {name}, ₽" for label in ("Без ИУ", "С ИУ") for _, name in COSTS]
    headers += ["Не хватает данных: без ИУ", "Не хватает данных: с ИУ", "Примечание к хранению"]
    sheet.append(headers)
    for item in rows:
        row = item.row
        turnover = row.calculator_turnover
        standard, current, delta = row.profit, row.current_profit, row.delta
        coverage = turnover.coverage
        values = [
            item.store_name,
            str(row.article),
            row.name,
            row.scheme,
            row.subject,
            ", ".join(row.managers or []),
            item.period_from,
            item.period_to,
            turnover.value,
            ("Да" if coverage.complete else "Нет") if coverage else None,
            standard.margin,
            standard.roi,
            current.margin,
            current.roi,
            delta.margin,
            delta.roi,
        ]
        values += [getattr(scenario, key) for scenario in (standard, current) for key, _ in COSTS]
        values += ["; ".join(scenario.missing or []) for scenario in (standard, current)]
        warning = row.storage_calculation.warning
        values.append(warning)
        sheet.append(values)
        if warning:
            storage_cell = sheet.cell(sheet.max_row, 27)
            storage_cell.comment = Comment(warning, "CheckStock")
            storage_cell.number_format = '#,##0" †"'
            storage_cell.fill = PatternFill("solid", fgColor="FFF3CD")
        for cell in sheet[sheet.max_row]:
            # Names and external data must never become executable Excel formulas.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif isinstance(cell.value, (int, float)):
                if not cell.comment:
                    cell.number_format = "#,##0"
        for column in (11, 12, 13, 14, 15, 16):
            cell = sheet.cell(sheet.max_row, column)
            loss = (standard if column < 13 else current).margin
            if column >= 15 and isinstance(cell.value, (int, float)) and cell.value:
                cell.font = Font(color="C62828" if cell.value < 0 else "188038")
            elif column < 15 and loss is not None and loss < 0:
                cell.font = Font(color="C62828")
                cell.fill = PatternFill("solid", fgColor="FDECEC")
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="32566F")
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22
    sheet.column_dimensions["C"].width = 42
    sheet.column_dimensions["B"].width = 18
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
