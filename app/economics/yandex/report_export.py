"""YM report workbooks with the same report rows, totals and daily detail as the UI."""

import io

SUMMARY = (
    ("name", "Товар"),
    ("article", "Артикул"),
    ("store_name", "Магазин"),
    ("subject", "Категория"),
    ("manager", "Менеджер"),
    ("orders_count", "Заказы, шт."),
    ("orders_amount", "ТО заказов, ₽"),
    ("cancel_count", "Отмены, шт."),
    ("cancel_amount", "Отмены, ₽"),
    ("net_orders_count", "Заказы − отмены, шт."),
    ("net_orders_amount", "ТО после отмен, ₽"),
    ("buyout_percent", "Выкуп, %"),
    ("stock", "Остаток всего"),
    ("stock_fbs", "FBS"),
    ("stock_fbo", "FBY"),
    ("stock_fulfillment", "ФФ"),
    ("stock_days", "Хватит, дней"),
    ("impressions", "Показы"),
    ("clicks", "Клики"),
    ("ctr", "CTR, %"),
    ("cpc", "CPC, ₽"),
    ("advertising_spend", "Реклама, ₽"),
    ("drr", "ДРР с выкупом, %"),
    ("margin_orders_count", "Расчётные выкупы"),
    ("margin", "Маржа периода, ₽"),
    ("purchase_value", "Закупка периода, ₽"),
    ("roi", "ROI, %"),
    ("margin_complete", "Полная история прибыли"),
    ("margin_missing_days", "Даты неполного расчёта"),
    ("status", "Полнота расчёта"), ("messages", "Не учтены / неизвестны: параметры и даты"),
    ("unavailable_days", "Дни без основы расчёта"),
    ("roi_purchase_value", "Закупка в базе ROI, ₽"),
)
DAILY = (
    ("date", "Дата"),
    ("orders_count", "Заказы, шт."),
    ("net_orders_count", "Заказы − отмены, шт."),
    ("advertising_spend", "Реклама, ₽"),
    ("buyout_percent", "Выкуп, %"),
    ("expected_buyouts", "Расчётные выкупы"),
    ("retail_price", "Цена без СПП, ₽"),
    ("customer_price", "Цена с СПП, ₽"),
    ("pay_price", "Цена с Пэй, ₽"),
    ("purchase_price", "Закупка на выкуп, ₽"),
    ("fulfillment_expense", "ФФ на выкуп, ₽"),
    ("logistics_delivery", "Доставка выкупа, ₽"),
    ("logistics_returns", "Обратная доставка невыкупов, ₽"),
    ("logistics_repeat_delivery", "Повторная доставка, ₽"),
    ("logistics_transit", "Транзит, ₽"),
    ("logistics_adjustment", "Корректировка логистики, ₽"),
    ("logistics", "Логистика итого, ₽"),
    ("commission_percent", "Комиссия ЯМ, %"),
    ("commission_value", "Комиссия ЯМ на выкуп, ₽"),
    ("acquiring_percent", "Перевод платежа (Эквайринг 1), %"),
    ("acquiring_value", "Перевод платежа (Эквайринг 1) на выкуп, ₽"),
    ("payment_acceptance", "Приём платежа (Эквайринг 2) на выкуп, ₽"),
    ("team_commission_percent", "Комиссия компании, %"),
    ("team_commission_value", "Комиссия компании на выкуп, ₽"),
    ("vat_percent", "НДС, %"),
    ("vat_value", "НДС на выкуп, ₽"),
    ("usn_percent", "УСН, %"),
    ("usn_value", "УСН на выкуп, ₽"),
    ("loss_percent", "Потери от закупки, %"),
    ("loss", "Потери от закупки, ₽"),
    ("disposal_cost", "Тариф утилизации, ₽"),
    ("disposal", "Утилизация на выкуп, ₽"),
    ("advertising_per_unit", "Реклама на выкуп, ₽"),
    ("total_cost", "Всего расходов на выкуп, ₽"),
    ("net_profit", "Прибыль на выкуп, ₽"),
    ("net_revenue", "Чистая выручка на выкуп, ₽"),
    ("day_profit", "Прибыль дня, ₽"),
    ("day_purchase_value", "Закупка дня, ₽"),
    ("available", "Расчёт доступен"),
    ("status", "Полнота расчёта"), ("messages", "Не учтены / неизвестны: параметры и даты"),
    ("calculation_version", "Версия расчёта"),
)
TARGET = (
    ("name", "Товар"),
    ("article", "Артикул"),
    ("store_name", "Магазин"),
    ("code", "Код"),
    ("current_price", "Текущая цена с Пэй, ₽"),
    ("current_drr", "Текущий ДРР, %"),
    ("current_roi", "Текущий ROI, %"),
    ("target_price", "Целевая цена с Пэй, ₽"),
    ("target_spp_price", "Целевая цена с СПП, ₽"),
    ("target_retail_price", "Целевая цена продавца, ₽"),
    ("target_drr", "Целевой ДРР, %"),
    ("target_roi", "Целевой ROI, %"),
    ("target_actual_roi", "Расчётный ROI, %"),
    ("target_warnings", "Примечания"),
)


def safe(value):
    if isinstance(value, (list, tuple)):
        value = "; ".join(str(item) for item in value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value


def sheet(workbook, title, columns, rows, period):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = workbook.create_sheet(title)
    ws.append(["Яндекс Маркет · " + title + " · " + period])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    ws.append([label for _, label in columns])
    for row in rows:
        ws.append([safe(row.get(key)) for key, _ in columns])
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFE177")
        cell.alignment = Alignment(wrap_text=True)
    for index, (_, label) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = (
            34 if index == 1 else min(max(len(label), 17), 35)
        )
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{max(ws.max_row, 2)}"
    return ws


def build_xlsx(report, *, target=False):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    period = report["period_from"] + " — " + report["period_to"]
    if target:
        sheet(workbook, "Целевая цена ЯМ", TARGET, report["rows"], period)
    else:
        rows = report.get("category_rows") if report.get("group_by") == "subject" else report["rows"]
        sheet(
            workbook, "Юниточная прибыль ЯМ", SUMMARY, [{"name": "Итого", **report["totals"]}, *rows], period
        )
        daily = [
            {**{key: product.get(key) for key in ("store_name", "article", "name")}, **day}
            for product in report["rows"]
            for day in product.get("daily_calculations", [])
        ]
        sheet(
            workbook,
            "Расчёт по дням",
            (("store_name", "Магазин"), ("article", "Артикул"), ("name", "Товар")) + DAILY,
            daily,
            period,
        )
        ws = workbook.create_sheet("Источники")
        ws.append(["Яндекс Маркет · данные из локальной БД"])
        ws.append(["Заказы и рекламные расходы за выбранные даты; текущие остатки."])
        ws.append(
            ["Прибыль и закупка за прошедшие дни — сохранённые дневные расчёты. Сегодня — текущие данные."]
        )
        ws.append(["ROI = прибыль периода / закупка тех же дней × 100. Пропуски не заменяются нулями."])
        ws.append(["ДРР = реклама / сумма (ТО товара × его выкуп / 100) × 100."])
        ws.column_dimensions["A"].width = 130
    data = io.BytesIO()
    workbook.save(data)
    prefix = "ym_target_price" if target else "ym_unit_profit"
    return data.getvalue(), f"{prefix}_{report['period_from']}_{report['period_to']}.xlsx"
